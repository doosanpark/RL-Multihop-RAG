"""PPT에 BC-only(지도학습) 베이스라인 결과와 주장 범위 한계를 반영.

기존 RL 발표자료(RL_Project_Presentation.pptx)에:
  - 슬라이드 16/17 결과표에 'BC-only (지도학습)' 행 추가 (행 높이 정규화로 전체 높이 보존)
  - 슬라이드 16/17 캡션에 BC≈RL 통찰 문장 추가
  - 슬라이드 30 종합결론 PART1 박스에 'BC 동급 → 한계는 정책 표현 용량' 추가
  - 슬라이드 31 한계 박스에 'bridge +0.01 사실상 미해결' 한계 추가
  - 슬라이드 4 한눈에 박스 PART1 줄에 BC 3점 비교 메시지 추가

서식 보존 원칙: 색상이 입혀진 multi-run 문단은 건드리지 않고, 마지막 run/문단을
deepcopy하여 텍스트만 바꿔 뒤에 덧붙인다.

실행:  python -m src.sol_a.add_bc_slides
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent.parent
PPTX = ROOT / "RL_Project_Presentation.pptx"


# ---------- 서식 보존 헬퍼 ----------

def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"shape {name!r} not found")


def table_of(slide):
    for sh in slide.shapes:
        if sh.has_table:
            return sh.table
    raise KeyError("no table")


def set_para_text_keep_first_run(para: _Paragraph, text: str) -> None:
    """문단의 첫 run 서식을 유지한 채 텍스트만 교체, 나머지 run 제거."""
    runs = para.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        para.text = text


def set_cell_text(cell, text: str) -> None:
    tf = cell.text_frame
    set_para_text_keep_first_run(tf.paragraphs[0], text)
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def append_run_like_last(para: _Paragraph, text: str) -> None:
    """문단 마지막 run을 복제해 텍스트만 바꿔 뒤에 추가 (색상 run 보존)."""
    runs = para.runs
    if not runs:
        para.add_run().text = text
        return
    last_r = runs[-1]._r
    new_r = deepcopy(last_r)
    t_el = new_r.find(qn("a:t"))
    if t_el is None:
        t_el = new_r.makeelement(qn("a:t"), {})
        new_r.append(t_el)
    t_el.text = text
    last_r.addnext(new_r)


def append_para_like(tf, template_para: _Paragraph, text: str) -> _Paragraph:
    """기존 문단을 복제해 본문 끝에 추가하고 텍스트 교체."""
    new_p = deepcopy(template_para._p)
    template_para._p.getparent().append(new_p)
    np = _Paragraph(new_p, tf)
    set_para_text_keep_first_run(np, text)
    return np


def add_table_row_after(table, src_row_idx: int, values: list[str]) -> None:
    """src_row를 복제해 그 바로 아래에 새 행 삽입 + 텍스트 채움.
    이후 모든 행 높이를 정규화해 표 전체 높이를 보존한다."""
    total_h = sum(r.height for r in table.rows)
    src_tr = table.rows[src_row_idx]._tr
    new_tr = deepcopy(src_tr)
    src_tr.addnext(new_tr)
    # 새 행에 값 채우기
    new_row = table.rows[src_row_idx + 1]
    for cell, val in zip(new_row.cells, values):
        set_cell_text(cell, val)
    # 행 높이 정규화 (전체 높이 유지)
    n = len(table.rows)
    each = int(total_h / n)
    for r in table.rows:
        r.height = Emu(each)


def main() -> None:
    prs = Presentation(str(PPTX))
    sl = list(prs.slides)

    # ----- 슬라이드 16: in-domain 표 + 캡션 -----
    s16 = sl[15]
    t16 = table_of(s16)
    # 행 순서: 0 header,1 Oracle,2 Naive,3 use_all,4 Step,5 Sparse,6 random
    add_table_row_after(t16, 5, ["BC-only (지도학습)", "0.335 ±.005", "0.546", "2.0"])
    cap16 = shape_by_name(s16, "Text 4")
    append_run_like_last(
        cap16.text_frame.paragraphs[0],
        " BC-only(0.335)도 RL과 동급 → 한계는 RL 방식이 아니라 정책 표현 용량.",
    )

    # ----- 슬라이드 17: transfer 표 + 캡션 -----
    s17 = sl[16]
    t17 = table_of(s17)
    add_table_row_after(t17, 5, ["BC-only (지도학습)", "0.265 ±.035", "0.466"])
    cap17 = shape_by_name(s17, "Text 7")
    append_run_like_last(
        cap17.text_frame.paragraphs[1],
        " BC-only도 −21%로 동일 하락 → OOD 취약성은 RL이 아닌 학습형 정책 공통의 성질.",
    )

    # ----- 슬라이드 30: 종합결론 PART1 박스 -----
    s30 = sl[29]
    box30 = shape_by_name(s30, "Text 8")  # PART1 설명 박스
    tf30 = box30.text_frame
    append_para_like(
        tf30, tf30.paragraphs[-1],
        "지도학습(BC)만으로도 RL과 동급(0.335 vs 0.355) → 한계는 RL 방식이 아니라 정책 표현 용량.",
    )

    # ----- 슬라이드 31: 한계 박스에 bridge 미해결 추가 -----
    s31 = sl[30]
    box31 = shape_by_name(s31, "Text 7")  # 한계 박스
    tf31 = box31.text_frame
    hdr_tmpl = tf31.paragraphs[2]   # '영어 단일 언어 ...' (header 스타일, 1-run)
    det_tmpl = tf31.paragraphs[3]   # '→ 전이 일반화 ...\n' (detail 스타일, 1-run)
    # 기존 마지막 항목(REINFORCE detail)은 trailing '\n'이 없어 새 항목이 바로 붙는다.
    # 다른 항목들처럼 빈 줄 간격을 주기 위해 '\n'을 덧붙인다.
    last_det = tf31.paragraphs[-1]
    if last_det.runs:
        last_det.runs[-1].text = last_det.runs[-1].text + "\n"
    append_para_like(tf31, hdr_tmpl, "RL 이득은 comparison에 집중")
    append_para_like(tf31, det_tmpl, "→ bridge(multi-hop 본질)는 +0.01로 사실상 미해결")

    # ----- 슬라이드 4: 한눈에 박스 PART1 줄 -----
    s4 = sl[3]
    box4 = shape_by_name(s4, "Text 15")  # 핵심 결과 박스
    append_run_like_last(
        box4.text_frame.paragraphs[0],
        "  ·  BC(지도학습) 0.335도 동급 → 병목은 정책 표현 용량",
    )

    prs.save(str(PPTX))
    print(f"[saved] {PPTX}")


if __name__ == "__main__":
    main()
